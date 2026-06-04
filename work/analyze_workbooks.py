import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


FILES = [
    Path(r"C:\Users\L03544646\Downloads\Indicadores .xlsx"),
    Path(r"C:\Users\L03544646\Downloads\Uniformes de Equipo RecSports 26.xlsx"),
]


def safe_value(value):
    if value is None:
        return None
    text = str(value).strip()
    if len(text) > 220:
        return text[:217] + "..."
    return text


def summarize_sheet(ws):
    formula_count = 0
    nonempty = 0
    formula_samples = []
    references = Counter()
    header_candidates = []
    text_terms = Counter()
    validations = []

    for dv in ws.data_validations.dataValidation:
        validations.append(
            {
                "type": dv.type,
                "operator": dv.operator,
                "formula1": safe_value(dv.formula1),
                "sqref": safe_value(dv.sqref),
            }
        )

    for r in range(1, min(ws.max_row, 80) + 1):
        vals = []
        count = 0
        for c in range(1, min(ws.max_column, 60) + 1):
            v = ws.cell(r, c).value
            vals.append(safe_value(v))
            if v not in (None, ""):
                count += 1
        if count >= 3 and len(header_candidates) < 8:
            header_candidates.append({"row": r, "values": vals[:30]})

    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v in (None, ""):
                continue
            nonempty += 1
            if isinstance(v, str):
                if v.startswith("="):
                    formula_count += 1
                    if len(formula_samples) < 20:
                        formula_samples.append({"cell": cell.coordinate, "formula": safe_value(v)})
                    for ref in re.findall(r"'?([^'!]+)'?![A-Z]{1,3}\$?\d+", v):
                        if ref and ref != ws.title:
                            references[ref] += 1
                else:
                    for token in re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]{4,}", v):
                        text_terms[token.lower()] += 1

    dims = {}
    for col in range(1, min(ws.max_column, 25) + 1):
        letter = openpyxl.utils.get_column_letter(col)
        vals = []
        for row in range(1, min(ws.max_row, 25) + 1):
            v = ws.cell(row, col).value
            if v not in (None, ""):
                vals.append(safe_value(v))
        if vals:
            dims[letter] = vals[:12]

    return {
        "sheet": ws.title,
        "rows": ws.max_row,
        "cols": ws.max_column,
        "nonempty": nonempty,
        "formulas": formula_count,
        "merged_ranges": len(ws.merged_cells.ranges),
        "tables": list(ws.tables.keys()),
        "validations": validations[:20],
        "header_candidates": header_candidates,
        "formula_samples": formula_samples,
        "external_sheet_refs": references.most_common(20),
        "top_terms": text_terms.most_common(30),
        "first_columns_preview": dims,
    }


def main():
    out = {}
    for file in FILES:
        wb = openpyxl.load_workbook(file, read_only=False, data_only=False)
        out[str(file)] = {
            "sheet_count": len(wb.sheetnames),
            "sheets": [summarize_sheet(ws) for ws in wb.worksheets],
        }
    Path("work/workbook_inventory.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    for file, data in out.items():
        print(str(file).encode("ascii", "ignore").decode("ascii"))
        print("sheets:", data["sheet_count"])
        for s in data["sheets"]:
            line = f"- {s['sheet']}: {s['rows']}x{s['cols']}, nonempty={s['nonempty']}, formulas={s['formulas']}, refs={s['external_sheet_refs'][:5]}"
            print(line.encode("ascii", "ignore").decode("ascii"))


if __name__ == "__main__":
    main()
