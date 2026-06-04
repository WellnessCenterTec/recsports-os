import json
from pathlib import Path

import openpyxl


SOURCE = Path(r"C:\Users\L03544646\Downloads\Uniformes de Equipo RecSports 26.xlsx")
OUT_JSON = Path("outputs/recsports-platform-prototype/uniformes-data.json")
OUT_SUMMARY = Path("work/uniformes_summary.json")


def value(cell):
    v = cell.value
    if v is None:
        return ""
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v).strip()


def sheet_rows(ws, max_rows=None):
    rows = []
    for row in ws.iter_rows():
        vals = [value(cell) for cell in row]
        if any(vals):
            rows.append(vals)
        if max_rows and len(rows) >= max_rows:
            break
    return rows


def normalize_table(rows):
    if not rows:
        return {"headers": [], "records": []}
    header_idx = 0
    best_count = 0
    for idx, row in enumerate(rows[:12]):
        count = sum(1 for item in row if item)
        if count > best_count:
            best_count = count
            header_idx = idx
    headers = rows[header_idx]
    headers = [h if h else f"Campo {i+1}" for i, h in enumerate(headers)]
    records = []
    for row in rows[header_idx + 1 :]:
        if not any(row):
            continue
        record = {}
        for i, header in enumerate(headers):
            if i < len(row):
                record[header] = row[i]
        records.append(record)
    return {"headers": headers, "records": records}


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    data = {}
    summary = {}
    for ws in wb.worksheets:
        rows = sheet_rows(ws)
        table = normalize_table(rows)
        data[ws.title] = table
        summary[ws.title] = {
            "rows": ws.max_row,
            "cols": ws.max_column,
            "headers": table["headers"],
            "sample_records": table["records"][:5],
            "record_count": len(table["records"]),
        }

    OUT_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    OUT_SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    for sheet, info in summary.items():
        print(f"{sheet}: {info['record_count']} records")
        print("headers:", info["headers"][:20])
        print("sample:", info["sample_records"][:1])


if __name__ == "__main__":
    main()
